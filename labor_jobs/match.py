"""Find occupations and job postings that fit a resume or a college major.

    python -m labor_jobs.match --resume data/resume.pdf --results output/adzuna-2026-10
    python -m labor_jobs.match --major "Mechanical Engineering" --state TX CA

Works from the files a main run wrote (postings_with_soc.csv,
detailed_occupations.csv), so run `python -m labor_jobs` on a postings file
first. A resume can be PDF, Word (.docx) or plain text. It is read in memory
only; outputs list what was found in it (job titles, technologies, degree,
major), never the resume text.

Each occupation gets a fit score from up to three signals:
  title  the resume lists a job title that maps to the occupation
  major  the college major leads to it, per the NCES/BLS CIP-SOC crosswalk
  skills technologies on the resume that the occupation's postings ask for
          (or O*NET lists for it)
Occupations that typically need a higher degree than the person has score
half.

Each posting in the best-fitting occupations is then checked against what it
states it requires:
  skills      share of the technologies its description names that the
              resume has
  experience  years it asks for ("5+ years of experience"), or what its title
              implies ("Senior" ~5), against years worked on the resume
  education   the degree it requires (allowing "or equivalent experience"), or
              the occupation's typical entry education, against the resume's
  occupation  the occupation fit above
Each check is meets, close, short or not stated. A posting whose stated
checks all meet (at least two of them) is a Strong match; two or more
shortfalls, or one with nothing met, is a Stretch; the rest are Partial
matches. This is a checklist of stated
requirements, not a prediction of being hired: no data here says who got
which job.
"""

import argparse
import os
import re
from pathlib import Path

import pandas as pd

from . import bls, h1b, majors, requirements, skills
from .matching import build_matcher
from .report import table
from .soc_mapping import load_onet_titles, load_title_overrides, normalize_title

WEIGHTS = {"title": 3.0, "major": 2.0, "skills": 2.0}
CONFIDENT_METHODS = {"exact", "override", "employer", "technology", "keyword", "partial"}
EDUCATION_CHOICES = {"none": 0, "associate": 2, "bachelor": 3, "master": 4, "doctoral": 5}


def read_document(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        import docx
        document = docx.Document(path)
        cells = [c.text for t in document.tables for row in t.rows for c in row.cells]
        return "\n".join([p.text for p in document.paragraphs] + cells)
    return path.read_text(errors="ignore")


def title_candidates(text, technologies=()):
    """Short phrases that could be job titles: lines and parts of lines
    split at separators, 1-7 words, without long numbers (dates, phones).
    Skill lists (lines with several commas) and bare technology names are
    left out."""
    tech = {normalize_title(t) for t in technologies}
    out = []
    for line in text.splitlines():
        if line.count(",") >= 3:
            continue
        for part in re.split(r"\s+[|–—-]\s+|\s+at\s+|,|\t|\s{3,}", line):
            part = part.strip(" •*·:")
            words = part.split()
            if 1 <= len(words) <= 7 and re.search(r"[A-Za-z]{3}", part) \
                    and not re.search(r"\d{3,}|@|www\.|https?:", part) \
                    and normalize_title(part) not in tech:
                out.append(part)
    return list(dict.fromkeys(out))


def resume_profile(text, matcher, technologies, major_matcher):
    """What the resume says, as plain data (no resume text kept)."""
    pattern, lookup = skills.mention_pattern(technologies)
    candidates = pd.Series(title_candidates(text, lookup), dtype=object)
    titles = {}
    if len(candidates):
        mapped = matcher.map_titles(candidates)
        for phrase, row in zip(candidates, mapped.itertuples()):
            if row.match_method in CONFIDENT_METHODS and isinstance(row.onet_code, str):
                titles.setdefault(row.onet_code, phrase)
    techs = sorted({lookup[m] for m in pattern.findall(text)})
    level, label = majors.education_level(text)
    return {"titles": titles, "technologies": techs,
            "majors": majors.majors_in_text(text), "education_level": level,
            "education": label, "years": requirements.resume_years(text)}


def _tech_names(value):
    """"Python 78%, SQL 35%" or "Python, SQL" -> {"Python", "SQL"}."""
    if not isinstance(value, str):
        return set()
    return {re.sub(r"\s+\d+%$", "", t.strip()) for t in value.split(",") if t.strip()}


def score_occupations(detailed, profile, major_socs):
    """Fit score 0-1 per occupation in the results, with the reasons."""
    held = profile["titles"]
    held_socs = {code[:7] for code in held}
    resume_tech = set(profile["technologies"])
    # Every signal the person provided counts for every occupation (as 0
    # where it doesn't match), so occupations aren't scored on fewer signals.
    available = [k for k, present in (("title", bool(held)), ("major", major_socs is not None),
                                      ("skills", bool(resume_tech))) if present]
    rows = []
    for occ in detailed.itertuples(index=False):
        signals, why = dict.fromkeys(available, 0.0), []
        if occ.onet_code in held:
            signals["title"] = 1.0
            why.append(f"held title: {held[occ.onet_code]}")
        elif occ.soc_code in held_socs:
            signals["title"] = 0.6
            why.append("held a title in the same occupation group")
        if major_socs is not None and occ.soc_code in major_socs:
            signals["major"] = 1.0
            why.append("fits the major")
        occ_tech = _tech_names(getattr(occ, "posting_technologies", None)) or \
            _tech_names(getattr(occ, "onet_hot_technologies", None))
        if resume_tech:
            overlap = sorted(resume_tech & occ_tech)
            signals["skills"] = min(1.0, len(overlap) / min(len(resume_tech), 5))
            if overlap:
                why.append("skills: " + ", ".join(overlap[:6]))
        if not available:
            continue
        fit = sum(WEIGHTS[k] * signals[k] for k in available) / sum(WEIGHTS[k] for k in available)
        needs = majors.BLS_EDUCATION.get(getattr(occ, "typical_education", None))
        if needs is not None and needs > profile["education_level"]:
            fit *= 0.5
            why.append(f"usually needs: {occ.typical_education}")
        missing = [t for t in (_tech_names(getattr(occ, "posting_technologies", None)) - resume_tech)]
        rows.append({"onet_code": occ.onet_code, "fit_score": round(fit, 3),
                     "why": "; ".join(why), "skills_to_add": ", ".join(sorted(missing)[:6])})
    scores = pd.DataFrame(rows, columns=["onet_code", "fit_score", "why", "skills_to_add"])
    out = detailed.merge(scores, on="onet_code")
    out = out[out["fit_score"] > 0]
    return out.sort_values(["fit_score", "postings"], ascending=False)


def rank_postings(postings, occupations, profile, top_occupations, states=None, metro=None,
                  metro_names=None, min_salary=None):
    best = occupations.head(top_occupations)
    fit = dict(zip(best["onet_code"], best["fit_score"]))
    names = dict(zip(best["onet_code"], best["onet_title"]))
    df = postings[postings["onet_code"].isin(fit)].copy()
    if states:
        df = df[df["state"].isin([s.upper() for s in states])]
    if metro and "metro" in df and metro_names is not None:
        wanted = {code for code, name in metro_names.items() if metro.lower() in str(name).lower()}
        df = df[df["metro"].astype(str).str.replace(r"\.0$", "", regex=True).isin(wanted)]
    if min_salary:
        df = df[df["annual_salary"] >= min_salary]
    resume_tech = set(profile["technologies"])
    posting_tech = df.get("technologies", pd.Series("", index=df.index)).fillna("") \
        .map(lambda s: {t.strip() for t in s.split(";") if t.strip()})
    df["matched_skills"] = posting_tech.map(lambda t: ", ".join(sorted(t & resume_tech)))
    df["missing_skills"] = posting_tech.map(lambda t: ", ".join(sorted(t - resume_tech)))
    coverage = posting_tech.map(lambda t: len(t & resume_tech) / len(t) if t else 0.0)
    df["occupation"] = df["onet_code"].map(names)
    df["match_score"] = (df["onet_code"].map(fit) + 0.5 * coverage).round(3)
    education = dict(zip(best["onet_code"], best.get("typical_education", pd.Series(dtype=object))))
    checks = [posting_checks(row, posting_tech[i], profile, fit.get(row.onet_code),
                             education.get(row.onet_code))
              for i, row in zip(df.index, df.itertuples(index=False))]
    df = df.join(pd.DataFrame(checks, index=df.index))
    df["rating_order"] = df["rating"].map(RATING_ORDER)
    return df.sort_values(["rating_order", "match_score", "annual_salary"],
                          ascending=[True, False, False]).drop(columns="rating_order")


RATINGS = ("Strong match", "Partial match", "Stretch")
RATING_ORDER = {r: i for i, r in enumerate(RATINGS)}


def _grade(have, need, close):
    if have >= need:
        return "meets"
    return "close" if have >= close else "short"


def posting_checks(row, posting_tech, profile, occupation_fit, occupation_education):
    """The four checks for one posting, their descriptions, and the rating."""
    resume_tech = set(profile["technologies"])
    out = {}
    # Skills the posting names.
    if posting_tech and profile.get("from_resume"):
        share = len(posting_tech & resume_tech) / len(posting_tech)
        out["skills_check"] = "meets" if share >= 0.6 else ("close" if share >= 0.3 else "short")
        out["skills_detail"] = f"{len(posting_tech & resume_tech)} of {len(posting_tech)} listed"
    else:
        out["skills_check"], out["skills_detail"] = "not stated", "no technologies listed"             if not posting_tech else "no resume"
    # Experience.
    need = getattr(row, "years_required", None)
    need = None if need is None or pd.isna(need) else float(need)
    basis = "asked"
    if need is None:
        need, label = requirements.seniority(row.title)
        basis = label.lower() if label else None
    have = profile.get("years")
    if need is None or have is None:
        out["experience_check"] = "not stated"
        out["experience_detail"] = "no years asked" if need is None else "no dates on resume"
    else:
        out["experience_check"] = _grade(have, need, max(need - 1, need * 0.75))
        asked = f"{need:g}+ yrs {basis}" if basis == "asked" else f"{basis} (~{need:g}+ yrs)"
        out["experience_detail"] = f"{asked}; resume shows {have:g}"
    # Education.
    level = getattr(row, "degree_required", None)
    flexible = bool(getattr(row, "degree_flexible", False)) and not pd.isna(getattr(row, "degree_flexible", False))
    source = "posting"
    if level is None or pd.isna(level):
        level, source = majors.BLS_EDUCATION.get(occupation_education), "typical"
    has = profile.get("education_level")
    if level is None or has is None:
        out["education_check"], out["education_detail"] = "not stated", "no degree stated"
    else:
        level = int(level)
        name = requirements.DEGREE_NAMES.get(level, occupation_education or "degree")
        check = "meets" if has >= level else ("close" if flexible or has == level - 1 else "short")
        out["education_check"] = check
        out["education_detail"] = (f"{name} required" + (" or equivalent experience" if flexible else "")
                                   if source == "posting" else f"{name} typical for the occupation")
    # Occupation fit.
    if occupation_fit is None:
        out["occupation_check"] = "not stated"
    else:
        out["occupation_check"] = "meets" if occupation_fit >= 0.5 else ("close" if occupation_fit >= 0.25 else "short")
    grades = [out[k] for k in ("skills_check", "experience_check", "education_check", "occupation_check")]
    meets, close, short = grades.count("meets"), grades.count("close"), grades.count("short")
    if short == 0 and close == 0 and meets >= 2:
        out["rating"] = "Strong match"
    elif short >= 2 or (short and not meets):
        out["rating"] = "Stretch"
    else:
        out["rating"] = "Partial match"
    names = {"skills_check": "skills", "experience_check": "experience", "education_check": "education",
             "occupation_check": "occupation"}
    gaps = [names[k] for k in names if out[k] in ("short", "close")]
    out["rating_reason"] = ("meets what it states" if not gaps else
                            "check " + ", ".join(gaps))
    return out


def skills_to_add(postings, limit=10):
    """Technologies most often missing from the resume among the postings
    that are a strong or partial match, in occupations the resume fits well."""
    near = postings[postings["rating"].isin(RATINGS[:2]) & (postings["occupation_check"] == "meets")]
    counts = near["missing_skills"].fillna("").str.split(", ").explode()
    counts = counts[counts != ""].value_counts().head(limit)
    return counts.rename_axis("technology").reset_index(name="postings")


def build_report(profile, major_text, major_hits, occupations, postings, results_dir, top):
    lines = ["# Job matches", ""]
    if profile.get("from_resume"):
        held = ", ".join(sorted(set(profile["titles"].values()))) or "none recognized"
        lines += ["## From the resume",
                  f"- Job titles recognized: {held}",
                  f"- Technologies: {', '.join(profile['technologies']) or 'none recognized'}",
                  f"- Highest degree: {profile['education']}",
                  "- Years of work experience (from job dates): "
                  + (f"{profile['years']:g}" if profile.get("years") is not None else "no dates found"),
                  f"- Majors found: {', '.join(profile['majors']) or 'none'}",
                  "",
                  "The resume was read in memory only; nothing from it is saved beyond this list.",
                  ""]
    if major_hits:
        lines += ["## Major", f"\"{major_text}\" was matched to: " +
                  "; ".join(f"{title} (CIP {code})" for code, title, _ in major_hits), ""]
    lines += [f"## Best-fitting occupations (from `{results_dir}`)",
              table(occupations.head(top), {
                  "onet_title": "Occupation", "fit_score": "Fit", "postings": "Postings",
                  "median_posted_salary": "Posted median $", "bls_median_wage": "BLS median $",
                  "proj_growth_pct": "Growth % (10 yr)", "typical_education": "Typical education",
                  "why": "Why", "skills_to_add": "Skills in postings not on resume"}),
              "",
              "Fit combines held job titles, the major's usual occupations and shared "
              "technologies (0 to 1). Occupations that usually need a higher degree count half.",
              "",
              "## Postings checked against the resume",
              "Each posting's stated requirements against the resume. This is a checklist, "
              "not a prediction of being hired: no data here says who got which job, and "
              "employers weigh things a resume and a posting don't show. Experience counts "
              "the date ranges of jobs on the resume; education uses the posting's stated "
              "degree, or the occupation's typical one when the posting states none.",
              "",
              ", ".join(f"{r}: {int((postings['rating'] == r).sum()):,}" for r in RATINGS)
              if "rating" in postings else "",
              "",
              table(postings.head(top * 2), {
                  "rating": "Rating", "title": "Title", "company_name": "Company",
                  "location": "Location", "annual_salary": "Salary $",
                  "matched_skills": "Your skills it lists", "missing_skills": "Skills it lists you lack",
                  "experience_detail": "Experience", "education_detail": "Education",
                  "rating_reason": "Look at", "url": "Link"}),
              ""]
    if "rating" in postings and len(postings):
        gaps = skills_to_add(postings)
        if len(gaps):
            lines += ["## Skills to add",
                      "Technologies most often listed by strong and partial matches, in occupations "
                      "the resume fits well, that the resume doesn't mention.",
                      "",
                      table(gaps, {"technology": "Technology", "postings": "Postings listing it"}),
                      ""]
    return "\n".join(lines)


def main(argv=None):
    p = argparse.ArgumentParser(prog="labor_jobs.match",
                                description="Find occupations and postings that fit a resume or major.")
    p.add_argument("--resume", help="resume file (.pdf, .docx or .txt)")
    p.add_argument("--major", help='college major, e.g. "Computer Science" or "BS in Nursing"')
    p.add_argument("--education", choices=list(EDUCATION_CHOICES),
                   help="highest degree, if not on the resume (default: bachelor with --major alone)")
    p.add_argument("--results", default="output", help="output folder of a main run (default: output)")
    p.add_argument("--state", nargs="+", help="only postings in these states, e.g. TX CA")
    p.add_argument("--metro", help='only postings in metro areas whose name contains this, e.g. "Austin"')
    p.add_argument("--min-salary", type=float, help="only postings advertising at least this much a year")
    p.add_argument("--top", type=int, default=15, help="occupations to show (default 15)")
    p.add_argument("--out", help="folder for the match files (default: <results>/matches)")
    p.add_argument("--cache", default="data/cache")
    p.add_argument("--onet-version", default="29_3")
    p.add_argument("--onet-dir")
    p.add_argument("--oews-year", type=int, default=2024)
    p.add_argument("--h1b-year", type=int, default=2024)
    p.add_argument("--skip-h1b", action="store_true", help="don't use H-1B job titles for matching")
    p.add_argument("--contact-email", default=os.environ.get("BLS_CONTACT_EMAIL"))
    args = p.parse_args(argv)
    if not (args.resume or args.major):
        p.error("give --resume, --major or both")

    results = Path(args.results)
    out = Path(args.out) if args.out else results / "matches"
    out.mkdir(parents=True, exist_ok=True)
    postings = pd.read_csv(results / "postings_with_soc.csv", dtype={"onet_code": str, "soc_code": str},
                           low_memory=False)
    detailed = pd.read_csv(results / "detailed_occupations.csv", dtype={"onet_code": str, "soc_code": str})
    metro_names = None
    if (results / "metros.csv").exists():
        m = pd.read_csv(results / "metros.csv", dtype={"cbsa": str})
        metro_names = dict(zip(m["cbsa"], m["metro_name"]))

    user_agent = f"labor-jobs/1.0 ({args.contact_email})" if args.contact_email else "labor-jobs/1.0"
    cip_csv, crosswalk = majors.cip_files(args.cache, user_agent)
    major_matcher = majors.MajorMatcher(cip_csv, crosswalk)

    profile = {"titles": {}, "technologies": [], "majors": [], "education_level": 3,
               "education": "Bachelor's degree (assumed)"}
    if args.resume:
        profiles, technologies = skills.load_profiles(args.onet_version, args.cache, user_agent,
                                                      args.onet_dir)
        oews_national = bls.read_oews(bls.oews_file(args.cache, args.oews_year, "nat", user_agent))
        employer_titles = None
        if not args.skip_h1b:
            employer_titles = h1b.reference_titles(
                h1b.read_h1b(h1b.h1b_files(args.cache, args.h1b_year, user_agent)))
        matcher = build_matcher(oews_national,
                                load_onet_titles(args.onet_version, args.cache, user_agent, args.onet_dir),
                                profiles, technologies, load_title_overrides(),
                                employer_titles=employer_titles)
        profile = resume_profile(read_document(args.resume), matcher, technologies, major_matcher)
        profile["from_resume"] = True
    if args.education:
        profile["education_level"] = EDUCATION_CHOICES[args.education]
        profile["education"] = args.education.title()

    major_text = args.major or (profile["majors"][0] if profile["majors"] else None)
    major_hits = major_matcher.match(major_text) if major_text else []
    major_socs = set(major_matcher.occupations_for([c for c, _, _ in major_hits])) if major_hits else None
    if major_text and not major_hits:
        print(f"Couldn't match the major \"{major_text}\" to a college program; ignoring it.")

    occupations = score_occupations(detailed, profile, major_socs)
    ranked = rank_postings(postings, occupations, profile, args.top, args.state, args.metro,
                           metro_names, args.min_salary)
    occupations.to_csv(out / "matching_occupations.csv", index=False)
    ranked.drop(columns=[c for c in ("description",) if c in ranked]).to_csv(
        out / "matching_postings.csv", index=False)
    (out / "matches.md").write_text(build_report(profile, major_text, major_hits, occupations,
                                                 ranked, results, args.top))
    print(f"{len(occupations):,} fitting occupations, {len(ranked):,} postings. "
          f"Wrote {out}/matches.md, matching_occupations.csv, matching_postings.csv")


if __name__ == "__main__":
    main()
