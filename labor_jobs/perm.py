"""PERM labor certifications (employer-sponsored green cards) from the
Department of Labor: the wage offered for a specific permanent job, with
its employer and worksite. Published yearly at
https://www.dol.gov/agencies/eta/foreign-labor/performance; the Q4 file
covers the whole federal fiscal year.

The application form changed in 2023, so fiscal year 2024 comes in two
files. Old-form filings carry an occupation (SOC) code; new-form filings
don't, so they are matched to occupations by job title like postings are.
As with H-1B data, only job, employer, worksite and wage columns are read.
"""

from pathlib import Path

import pandas as pd
import requests

from .fetch import download
from .h1b import PER_YEAR, _slim
from .states import ABBRS, NAME_TO_ABBR

URL = "https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/{name}"
FILES = ("PERM_Disclosure_Data_FY{year}_Q4.xlsx", "PERM_Disclosure_Data_New_Form_FY{year}_Q4.xlsx")
OLD = ["CASE_NUMBER", "CASE_STATUS", "DECISION_DATE", "EMPLOYER_NAME", "JOB_TITLE", "PW_SOC_CODE",
       "WAGE_OFFER_FROM", "WAGE_OFFER_UNIT_OF_PAY", "WORKSITE_CITY", "WORKSITE_STATE"]
NEW = ["CASE_NUMBER", "CASE_STATUS", "DECISION_DATE", "EMP_BUSINESS_NAME", "JOB_TITLE",
       "JOB_OPP_WAGE_FROM", "JOB_OPP_WAGE_PER", "PRIMARY_WORKSITE_CITY", "PRIMARY_WORKSITE_STATE"]
RENAME = {"EMP_BUSINESS_NAME": "EMPLOYER_NAME", "JOB_OPP_WAGE_FROM": "WAGE_OFFER_FROM",
          "JOB_OPP_WAGE_PER": "WAGE_OFFER_UNIT_OF_PAY", "PRIMARY_WORKSITE_CITY": "WORKSITE_CITY",
          "PRIMARY_WORKSITE_STATE": "WORKSITE_STATE"}


def perm_files(cache_dir, fiscal_year, user_agent):
    """The fiscal year's files that exist (newer years have only one)."""
    paths = []
    for pattern in FILES:
        name = pattern.format(year=fiscal_year)
        try:
            paths.append(download(URL.format(name=name), Path(cache_dir) / name, user_agent))
        except requests.HTTPError:
            continue
    return paths


def _state(value):
    value = str(value or "").strip()
    return value.upper() if value.upper() in ABBRS else NAME_TO_ABBR.get(value.lower())


def read_perm(paths):
    """Certified filings with an annual offered wage, in the same layout as
    h1b.read_h1b. onet_code is empty for new-form filings."""
    frames = []
    for path in paths:
        path = Path(path)
        if path.suffix in (".csv", ".gz"):
            frame = pd.read_csv(path, dtype=str)
        else:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            header = {str(h).strip().upper() for h in next(wb[wb.sheetnames[0]].iter_rows(values_only=True)) if h}
            wb.close()
            frame = pd.read_csv(_slim(path, OLD if "PW_SOC_CODE" in header else NEW), dtype=str)
        frames.append(frame.rename(columns=RENAME))
    df = pd.concat(frames, ignore_index=True).drop_duplicates("CASE_NUMBER")
    df = df[df["CASE_STATUS"].str.upper().str.startswith("CERTIFIED")]

    code = df.get("PW_SOC_CODE", pd.Series(None, index=df.index, dtype=object)) \
        .str.strip().str.extract(r"^(\d{2}-\d{4})(\.\d{2})?")
    out = pd.DataFrame({
        "case_number": df["CASE_NUMBER"],
        "job_title": df["JOB_TITLE"].str.strip(),
        "onet_code": code[0] + code[1].fillna(".00"),
        "employer": df["EMPLOYER_NAME"].str.strip().str.upper(),
        "state": df["WORKSITE_STATE"].map(_state),
        "city": df["WORKSITE_CITY"].str.strip().str.title(),
        "decision_date": pd.to_datetime(df["DECISION_DATE"], errors="coerce"),
    })
    wage = pd.to_numeric(df["WAGE_OFFER_FROM"].str.replace(r"[$,]", "", regex=True), errors="coerce")
    out["annual_wage"] = wage * df["WAGE_OFFER_UNIT_OF_PAY"].str.upper().str.strip().map(PER_YEAR)
    out["annual_wage"] = out["annual_wage"].where(out["annual_wage"].between(15_000, 2_000_000))
    out = out.dropna(subset=["annual_wage"])
    out["soc_code"] = out["onet_code"].str[:7]
    return out.reset_index(drop=True)
